import importlib.util
import json
import tomllib
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[2]
FX = ROOT / "fixtures" / "northwind"

_spec = importlib.util.spec_from_file_location("build_model_xlsx", FX / "build_model_xlsx.py")
mod = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(mod)

EXPECTED_IDS = {
    "bs_imbalance",
    "hardcode_calc",
    "broken_link",
    "inconsistent_formula",
    "wc_sign",
}


def _committed_manifest():
    return json.loads((FX / "model_bugs.json").read_text(encoding="utf-8"))


def test_openpyxl_is_a_runtime_dependency():
    deps = tomllib.loads((ROOT / "pyproject.toml").read_text(encoding="utf-8"))["project"]["dependencies"]
    assert any(d.replace(" ", "").startswith("openpyxl") for d in deps)


def test_manifest_matches_rebuilt_workbook(tmp_path):
    xlsx = mod.build(tmp_path)
    rebuilt = json.loads((tmp_path / "model_bugs.json").read_text(encoding="utf-8"))
    committed = _committed_manifest()
    assert rebuilt == committed
    wb = openpyxl.load_workbook(xlsx)
    for b in committed["bugs"]:
        assert b["sheet"] in wb.sheetnames, b["id"]
        assert wb[b["sheet"]][b["cell"]].value is not None, b["id"]
    # recomputed amounts
    bs = wb["BS"]
    by_id = {b["id"]: b for b in committed["bugs"]}
    assert by_id["bs_imbalance"]["amount"] == round(bs["D8"].value - mod.CLEAN_OTHER_ASSETS_FY26, 6)
    assert by_id["wc_sign"]["amount"] == 2 * (bs["E6"].value - bs["D6"].value)


def test_every_bug_has_lowercase_match_phrases():
    for b in _committed_manifest()["bugs"]:
        assert b["match"], b["id"]
        assert all(m == m.lower() for m in b["match"]), b["id"]


def test_each_planted_bug_is_present(tmp_path):
    xlsx = mod.build(tmp_path)
    wb = openpyxl.load_workbook(xlsx)  # formulas, not cached values
    bugs = {b["id"]: b for b in _committed_manifest()["bugs"]}
    assert set(bugs) == EXPECTED_IDS

    b = bugs["bs_imbalance"]
    assert (b["sheet"], b["cell"]) == ("BS", "D8")
    assert wb["BS"]["D8"].value == 1012.5
    assert round(wb["BS"]["D8"].value - mod.CLEAN_OTHER_ASSETS_FY26, 6) == b["amount"] == 12.5
    assert wb["BS"]["D15"].value == "=D9-D14"

    b = bugs["hardcode_calc"]
    assert (b["sheet"], b["cell"]) == ("IS", "E8")
    assert wb["IS"]["E8"].value == b["value"] == 350.0
    assert not str(wb["IS"]["E8"].value).startswith("=")
    assert wb["IS"]["D8"].value == "=D5*Inputs!D6"

    b = bugs["broken_link"]
    assert (b["sheet"], b["cell"]) == ("IS", "D15")
    assert "#REF!" in str(wb["IS"]["D15"].value)

    b = bugs["inconsistent_formula"]
    assert (b["sheet"], b["cell"]) == ("IS", "E6")
    assert wb["IS"]["E6"].value == b["formula"]
    assert wb["IS"]["D6"].value == b["neighbour_formula"] == "=D5*(1-Inputs!D5)"
    assert wb["IS"]["E6"].value != "=E5*(1-Inputs!E5)"

    b = bugs["wc_sign"]
    assert (b["sheet"], b["cell"]) == ("CF", "E7")
    assert wb["CF"]["E7"].value == b["formula"]
    assert b["correct_formula"] == "=-" + b["formula"][1:]
    assert b["amount"] == 2 * (wb["BS"]["E6"].value - wb["BS"]["D6"].value) == 20.0
