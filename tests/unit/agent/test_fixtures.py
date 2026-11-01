import json
import re
from pathlib import Path

import openpyxl
import pytest

from skillhub.agent.fixtures import MAX_CHARS, FixtureStore, read_workbook_dump
from skillhub.evalspec import EvalCase


@pytest.fixture
def fx(tmp_path: Path) -> Path:
    root = tmp_path / "fixtures"
    (root / "northwind").mkdir(parents=True)
    (root / "northwind" / "consensus_q3.json").write_text(
        json.dumps({"ticker": "NVDA", "eps": 1.2, "source": "FactSet", "date": "2026-10-20"})
    )
    (root / "northwind" / "press_release.md").write_text("Northwind Robotics reports record Q3 revenue.")
    (root / "northwind" / "overnight_news.md").write_text("Analysts raise NVDA targets after supply update.")
    (root / "northwind" / "web_index.json").write_text(
        json.dumps(
            {
                "keywords": {
                    "nvda": ["northwind/overnight_news.md"],
                    "targets": ["northwind/overnight_news.md"],
                    "consensus": ["northwind/consensus_q3.json"],
                },
                "max_results": 3,
            }
        )
    )
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Model"
    ws["A1"] = "Revenue"
    ws["B1"] = 100
    ws["C1"] = "=B1*1.1"
    wb.save(root / "northwind" / "model.xlsx")
    (root.parent / "secret.txt").write_text("top secret")
    return root


def case(fixtures: list[str]) -> EvalCase:
    return EvalCase(id=1, prompt="preview NVDA", should_trigger=True, fixtures=fixtures)


def test_no_fixtures_means_no_tools(fx):
    assert FixtureStore(fx).tools_for(case([])) == []


def test_fixture_case_offers_three_fake_tools(fx):
    names = [t.name for t in FixtureStore(fx).tools_for(case(["northwind/consensus_q3.json"]))]
    assert names == ["web_search", "read_file", "read_workbook"]


def test_read_tools_name_the_case_fixture_paths(fx):
    """Seen live: the model was never told the paths, guessed `model.xlsx`, and every read was refused."""
    tools = {
        t.name: t
        for t in FixtureStore(fx).tools_for(case(["northwind/model.xlsx", "northwind/press_release.md"]))
    }
    for name in ("read_file", "read_workbook"):
        path = tools[name].parameters["properties"]["path"]
        assert path["enum"] == ["northwind/model.xlsx", "northwind/press_release.md"]
        assert "northwind/model.xlsx" in tools[name].description


@pytest.mark.parametrize(
    "asked",
    [
        "northwind/model.xlsx",
        "fixtures/northwind/model.xlsx",
        "Northwind/Model.xlsx",
        "model.xlsx",
        "./northwind/model.xlsx",
    ],
)
def test_read_workbook_accepts_harmless_variants_of_a_declared_path(fx, asked):
    out = FixtureStore(fx).call("read_workbook", {"path": asked}, case(["northwind/model.xlsx"]))
    assert "Revenue" in out and not out.startswith("Error")


@pytest.mark.parametrize("asked", ["secret.txt", "../secret.txt", "northwind/press_release.md"])
def test_paths_the_case_does_not_declare_are_still_refused(fx, asked):
    out = FixtureStore(fx).call("read_file", {"path": asked}, case(["northwind/model.xlsx"]))
    assert out.startswith("Error: not a fixture of this case") and "northwind/model.xlsx" in out


def test_an_ambiguous_bare_file_name_is_refused(fx, tmp_path):
    (fx / "other").mkdir()
    (fx / "other" / "model.xlsx").write_bytes((fx / "northwind" / "model.xlsx").read_bytes())
    out = FixtureStore(fx).call(
        "read_workbook", {"path": "model.xlsx"}, case(["northwind/model.xlsx", "other/model.xlsx"])
    )
    assert out.startswith("Error: not a fixture of this case")


def test_web_search_keyword_matches_fixture_documents(fx):
    out = FixtureStore(fx).call(
        "web_search", {"query": "NVDA targets"}, case(["northwind/overnight_news.md"])
    )
    assert out == "=== northwind/overnight_news.md ===\nAnalysts raise NVDA targets after supply update."


def test_web_search_no_match_returns_empty_marker(fx):
    out = FixtureStore(fx).call("web_search", {"query": "zebra"}, case(["northwind/overnight_news.md"]))
    assert out == "No results."


def test_web_search_only_returns_docs_in_case_fixtures(fx):
    out = FixtureStore(fx).call(
        "web_search", {"query": "consensus nvda"}, case(["northwind/consensus_q3.json"])
    )
    assert "overnight_news" not in out
    assert out.startswith("=== northwind/consensus_q3.json ===")


def test_web_search_ranks_by_distinct_keywords_then_path_and_caps_at_three(fx):
    docs = ["northwind/a.md", "northwind/b.md", "northwind/c.md", "northwind/d.md"]
    for d in docs:
        (fx / d).write_text(f"text of {d}")
    (fx / "northwind" / "web_index.json").write_text(
        json.dumps({"keywords": {"nvda": docs, "targets": ["northwind/b.md"]}, "max_results": 3})
    )
    out = FixtureStore(fx).call("web_search", {"query": "NVDA targets"}, case(docs))
    headers = re.findall(r"^=== (.+) ===$", out, flags=re.M)
    assert headers == ["northwind/b.md", "northwind/a.md", "northwind/c.md"]


def test_fixture_root_defaults_to_skillhub_fixtures_env(fx, monkeypatch):
    monkeypatch.setenv("SKILLHUB_FIXTURES", str(fx))
    assert FixtureStore().root == fx.resolve()
    assert FixtureStore(fx / "northwind").root == (fx / "northwind").resolve()  # an explicit root wins


def test_fixture_root_defaults_to_fixtures_without_env(monkeypatch):
    monkeypatch.delenv("SKILLHUB_FIXTURES", raising=False)
    assert FixtureStore().root == Path("fixtures").resolve()


def test_read_file_returns_fixture_text(fx):
    out = FixtureStore(fx).call(
        "read_file", {"path": "northwind/press_release.md"}, case(["northwind/press_release.md"])
    )
    assert out == "Northwind Robotics reports record Q3 revenue."


@pytest.mark.parametrize("bad", ["../secret.txt", "/etc/passwd", "northwind/../../secret.txt"])
def test_read_file_path_traversal_returns_error_string(fx, bad):
    out = FixtureStore(fx).call("read_file", {"path": bad}, case(["northwind/press_release.md", bad]))
    assert out.startswith("Error:")
    assert "top secret" not in out


def test_read_file_outside_case_fixtures_returns_error(fx):
    out = FixtureStore(fx).call(
        "read_file", {"path": "northwind/overnight_news.md"}, case(["northwind/press_release.md"])
    )
    assert out.startswith("Error:")
    assert "Analysts" not in out


def test_read_workbook_outside_case_fixtures_returns_error(fx):
    out = FixtureStore(fx).call("read_workbook", {"path": "northwind/model.xlsx"}, case([]))
    assert out.startswith("Error:")


def test_missing_file_and_unknown_tool_and_missing_arg_return_error(fx):
    store, c = FixtureStore(fx), case(["northwind/nope.md"])
    assert store.call("read_file", {"path": "northwind/nope.md"}, c).startswith("Error:")
    assert store.call("nope", {}, c).startswith("Error:")
    assert store.call("read_file", {}, c).startswith("Error:")


def test_read_workbook_gives_formulas_and_values_cell_by_cell(fx):
    out = FixtureStore(fx).call(
        "read_workbook", {"path": "northwind/model.xlsx"}, case(["northwind/model.xlsx"])
    )
    assert "Model!C1" in out
    assert "=B1*1.1" in out
    assert "100" in out


def test_read_workbook_dump_is_formula_and_value_per_cell(fx):
    lines = read_workbook_dump(fx / "northwind" / "model.xlsx").splitlines()
    assert any(ln.startswith("Model!B1") and "100" in ln for ln in lines)
    # no cached value is stored for a freshly written workbook
    assert any(ln.startswith("Model!C1") and "formula==B1*1.1" in ln and "value=None" in ln for ln in lines)


def test_web_search_renders_workbook_and_markdown_hits(fx):
    (fx / "northwind" / "web_index.json").write_text(
        json.dumps(
            {"keywords": {"model": ["northwind/model.xlsx", "northwind/press_release.md"]}, "max_results": 3}
        )
    )
    docs = ["northwind/model.xlsx", "northwind/press_release.md"]
    out = FixtureStore(fx).call("web_search", {"query": "model"}, case(docs))
    assert "=== northwind/model.xlsx ===\nModel!A1 value=Revenue" in out
    assert "=== northwind/press_release.md ===\nNorthwind Robotics" in out


def test_web_search_missing_indexed_doc_does_not_break_others(fx):
    (fx / "northwind" / "web_index.json").write_text(
        json.dumps(
            {"keywords": {"q3": ["northwind/gone.md", "northwind/press_release.md"]}, "max_results": 3}
        )
    )
    out = FixtureStore(fx).call(
        "web_search", {"query": "q3"}, case(["northwind/gone.md", "northwind/press_release.md"])
    )
    assert "=== northwind/gone.md ===\n[unreadable:" in out
    assert "Northwind Robotics reports" in out


def test_web_search_undecodable_doc_is_marked_unreadable(fx):
    (fx / "northwind" / "bin.md").write_bytes(b"\xff\xfe\x00bad")
    (fx / "northwind" / "web_index.json").write_text(json.dumps({"keywords": {"q3": ["northwind/bin.md"]}}))
    out = FixtureStore(fx).call("web_search", {"query": "q3"}, case(["northwind/bin.md"]))
    assert out.startswith("=== northwind/bin.md ===\n[unreadable:")


def test_read_file_truncates_with_marker(fx):
    (fx / "northwind" / "big.md").write_text("x" * (MAX_CHARS + 50))
    out = FixtureStore(fx).call("read_file", {"path": "northwind/big.md"}, case(["northwind/big.md"]))
    assert out == "x" * MAX_CHARS + f"\n[truncated: {MAX_CHARS} of {MAX_CHARS + 50} chars]"


def test_read_workbook_truncates_with_marker(fx):
    wb = openpyxl.Workbook()
    ws = wb.active
    for i in range(1, 2001):
        ws[f"A{i}"] = "some long cell text for padding"
    wb.save(fx / "northwind" / "big.xlsx")
    out = FixtureStore(fx).call("read_workbook", {"path": "northwind/big.xlsx"}, case(["northwind/big.xlsx"]))
    assert re.search(rf"\n\[truncated: {MAX_CHARS} of \d+ chars\]$", out)


def test_web_search_truncates_each_block(fx):
    (fx / "northwind" / "big.md").write_text("y" * (MAX_CHARS + 10))
    (fx / "northwind" / "web_index.json").write_text(json.dumps({"keywords": {"big": ["northwind/big.md"]}}))
    out = FixtureStore(fx).call("web_search", {"query": "big"}, case(["northwind/big.md"]))
    assert out.endswith(f"[truncated: {MAX_CHARS} of {MAX_CHARS + 10} chars]")


def test_symlink_escaping_root_is_rejected(fx):
    (fx / "northwind" / "link.md").symlink_to(fx.parent / "secret.txt")
    store, c = FixtureStore(fx), case(["northwind/link.md"])
    for tool in ("read_file", "read_workbook"):
        out = store.call(tool, {"path": "northwind/link.md"}, c)
        assert out.startswith("Error:")
        assert "top secret" not in out


def test_web_search_symlink_escape_is_unreadable_not_leaked(fx):
    (fx / "northwind" / "link.md").symlink_to(fx.parent / "secret.txt")
    (fx / "northwind" / "web_index.json").write_text(json.dumps({"keywords": {"q3": ["northwind/link.md"]}}))
    out = FixtureStore(fx).call("web_search", {"query": "q3"}, case(["northwind/link.md"]))
    assert "[unreadable:" in out
    assert "top secret" not in out


def test_malformed_web_index_returns_specific_error(fx):
    (fx / "northwind" / "web_index.json").write_text("{not json")
    out = FixtureStore(fx).call("web_search", {"query": "nvda"}, case(["northwind/overnight_news.md"]))
    assert out == "Error: invalid web_index.json (northwind/web_index.json)"


def test_the_payload_of_a_case_is_its_files_as_the_tools_show_them_and_its_part_of_the_index():
    from pathlib import Path

    root = Path(__file__).resolve().parents[3] / "fixtures"
    store = FixtureStore(root)
    case = EvalCase(id=1, prompt="p", should_trigger=True,
                    fixtures=["northwind/model.xlsx", "northwind/press_release_q3.md"])  # fmt: skip
    got = store.payload(case)
    assert got["files"] == {
        "northwind/model.xlsx": store.call("read_workbook", {"path": "northwind/model.xlsx"}, case),
        "northwind/press_release_q3.md": store.call(
            "read_file", {"path": "northwind/press_release_q3.md"}, case
        ),
    }
    assert got["search_index"] and all(
        set(docs) <= set(case.fixtures) for docs in got["search_index"].values()
    )
    assert store.payload(EvalCase(id=2, prompt="p", should_trigger=False)) == {
        "files": {},
        "search_index": {},
    }
