"""Build the Northwind Robotics model fixture and its planted-bug manifest.

Synthetic data for a fictional company. Run:
    uv run python fixtures/northwind/build_model_xlsx.py fixtures/northwind
"""

from __future__ import annotations

import json
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

BLUE = Font(color="0000FF")  # hard-coded input
BLACK = Font(color="000000")  # formula

CLEAN_OTHER_ASSETS_FY26 = 1000.0  # what the FY26E roll-forward needs in BS!D8

# "match" holds lowercase phrases the planted_bugs_found scorer looks for (R2-16).
BUGS = [
    {
        "id": "bs_imbalance",
        "match": ["balance sheet does not balance", "out of balance", "bs!d8", "12.5"],
        "sheet": "BS",
        "cell": "D8",
        "kind": "balance_sheet_imbalance",
        "amount": 12.5,
        "description": "FY26E other assets is typed as 1,012.5; the roll-forward needs 1,000.0. The FY26E check at BS!D15 reads 12.5.",
    },
    {
        "id": "hardcode_calc",
        "match": ["hard-coded", "hardcoded", "hard coded", "is!e8"],
        "sheet": "IS",
        "cell": "E8",
        "kind": "hardcoded_in_calc_row",
        "value": 350.0,
        "correct_formula": "=E5*Inputs!E6",
        "description": "FY27E SG&A is a typed number in a calculation row. The formula gives 344.4.",
    },
    {
        "id": "broken_link",
        "match": ["#ref", "broken link", "broken reference", "is!d15"],
        "sheet": "IS",
        "cell": "D15",
        "kind": "broken_cross_sheet_link",
        "formula": "=Inputs!#REF!",
        "description": "Memo cell refers to a deleted cell on the Inputs sheet. Nothing depends on it.",
    },
    {
        "id": "inconsistent_formula",
        "match": ["inconsistent formula", "is!e6", "breaks the pattern"],
        "sheet": "IS",
        "cell": "E6",
        "kind": "inconsistent_formula",
        "formula": "=E5*Inputs!E5",
        "neighbour_cell": "D6",
        "neighbour_formula": "=D5*(1-Inputs!D5)",
        "description": "FY27E COGS multiplies by the gross margin instead of one minus the gross margin. The rest of row 6 uses the D6 pattern.",
    },
    {
        "id": "wc_sign",
        "match": ["wrong sign", "sign error", "cf!e7", "working capital"],
        "sheet": "CF",
        "cell": "E7",
        "kind": "wrong_sign_working_capital",
        "formula": "=(BS!E6-BS!D6)",
        "correct_formula": "=-(BS!E6-BS!D6)",
        "amount": 20.0,
        "description": "FY27E change in receivables has the wrong sign. An increase in receivables should reduce cash. Cash is overstated by 20.0, so the FY27E check at BS!E15 reads 20.0.",
    },
]

MANIFEST = {
    "workbook": "model.xlsx",
    "note": "Synthetic test data. Northwind Robotics is a fictional company. Every planted defect is listed; all other formulas are intended to be correct.",
    "bugs": BUGS,
}


def _put(ws, addr: str, value) -> None:
    ws[addr] = value
    is_formula = isinstance(value, str) and value.startswith("=")
    ws[addr].font = BLACK if is_formula else BLUE


def _years(ws, cols: str, labels: list[str]) -> None:
    for col, label in zip(cols, labels, strict=True):
        ws[f"{col}3"] = label


def _build_inputs(wb: Workbook) -> None:
    ws = wb.active
    ws.title = "Inputs"
    ws["B1"] = "Northwind Robotics synthetic model, $M unless stated. Synthetic test data."
    _years(ws, "CDE", ["FY25A", "FY26E", "FY27E"])
    rows = [
        (4, "Revenue growth", 0.04, 0.035),
        (5, "Gross margin", 0.41, 0.415),
        (6, "SG&A % of revenue", 0.20, 0.20),
        (7, "Tax rate", 0.21, 0.21),
        (8, "D&A", 40.0, 42.0),
        (9, "Dividends", 20.0, 21.0),
        (10, "Capex", 55.0, 58.0),
    ]
    for r, label, d, e in rows:
        ws[f"B{r}"] = label
        _put(ws, f"D{r}", d)
        _put(ws, f"E{r}", e)


def _build_is(wb: Workbook) -> None:
    ws = wb["IS"]
    ws["B1"] = "Income statement ($M)"
    _years(ws, "CDE", ["FY25A", "FY26E", "FY27E"])
    ws["B5"] = "Revenue"
    _put(ws, "C5", 1600.0)
    _put(ws, "D5", "=C5*(1+Inputs!D4)")
    _put(ws, "E5", "=D5*(1+Inputs!E4)")
    ws["B6"] = "COGS"
    _put(ws, "C6", 944.0)
    _put(ws, "D6", "=D5*(1-Inputs!D5)")
    _put(ws, "E6", "=E5*Inputs!E5")  # planted: inconsistent formula
    ws["B7"] = "Gross profit"
    ws["B8"] = "SG&A"
    _put(ws, "C8", 320.0)
    _put(ws, "D8", "=D5*Inputs!D6")
    _put(ws, "E8", 350.0)  # planted: hard-coded number in a calculation row
    ws["E8"].font = BLACK  # looks like a formula cell, as a pasted value would
    ws["B9"] = "EBITDA"
    ws["B10"] = "D&A"
    _put(ws, "C10", 38.0)
    _put(ws, "D10", "=Inputs!D8")
    _put(ws, "E10", "=Inputs!E8")
    ws["B11"] = "EBIT"
    ws["B12"] = "Tax"
    _put(ws, "C12", 60.0)
    _put(ws, "D12", "=D11*Inputs!D7")
    _put(ws, "E12", "=E11*Inputs!E7")
    ws["B13"] = "Net income"
    for col in "CDE":
        _put(ws, f"{col}7", f"={col}5-{col}6")
        _put(ws, f"{col}9", f"={col}7-{col}8")
        _put(ws, f"{col}11", f"={col}9-{col}10")
        _put(ws, f"{col}13", f"={col}11-{col}12")
    ws["B15"] = "Memo: FY26 revenue guidance low"
    _put(ws, "D15", "=Inputs!#REF!")  # planted: broken cross-sheet link


def _build_bs(wb: Workbook) -> None:
    ws = wb["BS"]
    ws["B1"] = "Balance sheet ($M)"
    _years(ws, "CDE", ["FY25A", "FY26E", "FY27E"])
    ws["B4"] = "Assets"
    ws["B5"] = "Cash"
    _put(ws, "C5", 300.0)
    _put(ws, "D5", "=CF!D12")
    _put(ws, "E5", "=CF!E12")
    ws["B6"] = "Accounts receivable"
    _put(ws, "C6", 240.0)
    _put(ws, "D6", 260.0)
    _put(ws, "E6", 270.0)
    ws["B7"] = "PP&E, net"
    _put(ws, "C7", 900.0)
    _put(ws, "D7", "=C7+Inputs!D10-IS!D10")
    _put(ws, "E7", "=D7+Inputs!E10-IS!E10")
    ws["B8"] = "Other assets"
    _put(ws, "C8", 1000.0)
    _put(ws, "D8", 1012.5)  # planted: balance sheet off by 12.5
    _put(ws, "E8", 1000.0)
    ws["B9"] = "Total assets"
    ws["B11"] = "Debt"
    ws["B12"] = "Common stock"
    ws["B13"] = "Retained earnings"
    _put(ws, "C13", 1640.0)
    _put(ws, "D13", "=C13+IS!D13-Inputs!D9")
    _put(ws, "E13", "=D13+IS!E13-Inputs!E9")
    ws["B14"] = "Total liabilities and equity"
    ws["B15"] = "Check: assets less L&E (should be 0)"
    for col in "CDE":
        _put(ws, f"{col}9", f"=SUM({col}5:{col}8)")
        _put(ws, f"{col}11", 500.0)
        _put(ws, f"{col}12", 300.0)
        _put(ws, f"{col}14", f"=SUM({col}11:{col}13)")
        _put(ws, f"{col}15", f"={col}9-{col}14")


def _build_cf(wb: Workbook) -> None:
    ws = wb["CF"]
    ws["B1"] = "Cash flow ($M)"
    _years(ws, "DE", ["FY26E", "FY27E"])
    ws["B5"] = "Net income"
    _put(ws, "D5", "=IS!D13")
    _put(ws, "E5", "=IS!E13")
    ws["B6"] = "D&A"
    _put(ws, "D6", "=IS!D10")
    _put(ws, "E6", "=IS!E10")
    ws["B7"] = "Change in receivables"
    _put(ws, "D7", "=-(BS!D6-BS!C6)")
    _put(ws, "E7", "=(BS!E6-BS!D6)")  # planted: wrong sign
    ws["B8"] = "Cash from operations"
    ws["B9"] = "Capex"
    _put(ws, "D9", "=-Inputs!D10")
    _put(ws, "E9", "=-Inputs!E10")
    ws["B10"] = "Dividends"
    _put(ws, "D10", "=-Inputs!D9")
    _put(ws, "E10", "=-Inputs!E9")
    ws["B11"] = "Net change in cash"
    ws["B12"] = "Ending cash"
    for col in "DE":
        _put(ws, f"{col}8", f"=SUM({col}5:{col}7)")
        _put(ws, f"{col}11", f"={col}8+{col}9+{col}10")
    _put(ws, "D12", "=BS!C5+D11")
    _put(ws, "E12", "=BS!D5+E11")


def _maybe_recalc(xlsx: Path) -> None:
    """Store cached formula results when LibreOffice is installed; optional, never required."""
    soffice = shutil.which("soffice")
    if not soffice:
        return
    with tempfile.TemporaryDirectory() as tmp:
        subprocess.run(
            [soffice, "--headless", "--calc", "--convert-to", "xlsx", "--outdir", tmp, str(xlsx)],
            check=True,
            capture_output=True,
        )
        shutil.copy(Path(tmp) / xlsx.name, xlsx)


def build(out_dir: Path) -> Path:
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    _build_inputs(wb)
    for name in ["IS", "BS", "CF"]:
        wb.create_sheet(name)
    _build_is(wb)
    _build_bs(wb)
    _build_cf(wb)
    xlsx = out_dir / "model.xlsx"
    wb.save(xlsx)
    _maybe_recalc(xlsx)
    (out_dir / "model_bugs.json").write_text(json.dumps(MANIFEST, indent=2) + "\n", encoding="utf-8")
    return xlsx


if __name__ == "__main__":
    target = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(__file__).resolve().parent
    print(build(target))
